from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "交付文档"
OUT_PATH = OUT_DIR / "WZ-724上位机工作交接文档_V1.0.docx"

NAVY = "123247"
BLUE = "177A9E"
CYAN = "31B6C8"
INK = "1F2933"
MUTED = "647580"
LIGHT_BLUE = "E8F2F6"
LIGHT_GRAY = "F3F5F7"
PALE_YELLOW = "FFF6D9"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
GRID = "B8C8D1"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, bold=None, italic=None, color=None,
                 ascii_font="Calibri", east_asia="Microsoft YaHei UI"):
    run.font.name = ascii_font
    run._element.get_or_add_rPr()
    rfonts = run._element.rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start),
                       ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], indent=TABLE_INDENT_DXA):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=9, color=MUTED)


def setup_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=0) + 1
    next_num = max(existing_num, default=0) + 1

    def add_definition(fmt, text, font=None):
        nonlocal next_abs, next_num
        abstract_id = next_abs
        num_id = next_num
        next_abs += 1
        next_num += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    return add_definition("bullet", "•", "Arial"), add_definition("decimal", "%1.")


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=INK)
    return p


def add_bullet(doc, text, bullet_num_id):
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_number(doc, text, decimal_num_id):
    p = doc.add_paragraph()
    apply_num(p, decimal_num_id)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, lines: Iterable[str]):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F7F8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=8.8, color="243746", ascii_font="Consolas",
                     east_asia="Microsoft YaHei UI")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, font_size=9.2, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_run_font(run, size=9.2, bold=True, color=NAVY)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "FAFBFC")
        for col_idx, value in enumerate(values):
            p = cells[col_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if col_idx > 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, size=font_size,
                         bold=(first_col_bold and col_idx == 0), color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("WZ-724 ROV 上位机 | 技术工作交接")
    set_run_font(run, size=8.5, bold=True, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C9D6DD")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("内部技术资料  |  第 ")
    set_run_font(run, size=8.5, color=MUTED)
    add_page_field(p)
    run = p.add_run(" 页")
    set_run_font(run, size=8.5, color=MUTED)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)
    bullet_id, decimal_id = setup_numbering(doc)

    # Cover: editorial_cover pattern, compact_reference_guide body.
    doc.add_paragraph().paragraph_format.space_after = Pt(52)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("技术工作交接手册")
    set_run_font(r, size=11, bold=True, color=CYAN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("WZ-724 ROV 上位机")
    set_run_font(r, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("开发、部署、运行、维护与故障排查")
    set_run_font(r, size=14, color=BLUE)

    add_callout(doc, "交接目标", "让接手人员能够独立完成工程打开、编译发布、现场连接、功能修改、日志分析和安全验证。")
    doc.add_paragraph().paragraph_format.space_after = Pt(55)

    cover_meta = [
        ("文档版本", "V1.0"),
        ("编制日期", "2026-09-03"),
        ("项目状态", "MAVLink实机链路已验证；摄像头自启动当前已关闭"),
        ("移交人", "________________"),
        ("接收人", "________________"),
    ]
    add_table(doc, ["项目", "内容"], cover_meta, [2160, 7200], 9.5, True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WZ-724 ROV 项目组 · 内部技术资料")
    set_run_font(r, size=9.5, color=MUTED)

    page_break(doc)

    add_heading(doc, "文档控制与快速接手", 1)
    add_body(doc, "本文件描述当前工作区中的 Qt 上位机、香橙派网络桥接与视频链路。飞控固件只记录上位机依赖的定制接口，不替代完整的固件变更说明。")
    add_heading(doc, "接手人首先要确认的五件事", 2)
    for item in (
        "确认源码目录完整，重点检查 wz-724.pro、thirdparty、models、styles 和 tools。",
        "Qt Creator 使用 Desktop Qt 5.14.2 MinGW 64-bit Kit，并重新运行 qmake。",
        "笔记本有线网卡保持 192.168.1.137/24，香橙派保持 192.168.1.101/24。",
        "首次实机测试必须保持飞控未解锁、推进器断电或拆桨，并确认急停可用。",
        "当前香橙派摄像头服务自启动已关闭；MAVLink服务不受影响。",
    ):
        add_number(doc, item, decimal_id)

    add_heading(doc, "当前交付物", 2)
    add_table(doc, ["交付项", "位置/说明"], [
        ("Qt工程", "项目根目录 / wz-724.pro"),
        ("Release构建", "build-wz-724-mavlink-release/release/wz-724.exe"),
        ("便携版", "wz-724-portable-v17-mavlink/"),
        ("便携压缩包", "wz-724-portable-v17-mavlink.zip"),
        ("协议测试", "tests/mavlink_manager_protocol_test.cpp"),
        ("链路探针", "tests/mavlink_udp_probe.py"),
        ("香橙派部署模板", "orangepi-deployment/"),
    ], [2600, 6760], 9.2, True)

    add_callout(doc, "版本管理风险", "当前目录不是 Git 仓库。接手后建议先完整备份，再初始化 Git 并提交一个“已验证基线”版本，避免后续无法追溯改动。", PALE_YELLOW, "8A6500")

    page_break(doc)

    add_heading(doc, "1. 系统架构", 1)
    add_heading(doc, "1.1 端到端链路", 2)
    add_code_block(doc, [
        "Pixhawk / ArduSub",
        "  |  UART / MAVLink / 57600 baud",
        "  v",
        "Orange Pi 192.168.1.101",
        "  |-- /dev/ttyS0 <-> socat <-> UDP 15001     [飞控数据]",
        "  `-- USB UVC -> V4L2 -> FFmpeg/MediaMTX     [视频数据]",
        "                    |",
        "                    | Ethernet / 电力载波",
        "                    v",
        "Laptop 192.168.1.137",
        "  |-- UDP 5555 -> MAVLinkManager -> 仪表/3D/状态/控制",
        "  `-- RTSP :554 -> FFmpegRtspPlayer -> 显示/录像",
    ])
    add_body(doc, "电力载波在本系统中是透明以太网传输介质，正常情况下不改变 Qt 程序、IP、端口或 MAVLink消息。")

    add_heading(doc, "1.2 分层职责", 2)
    add_table(doc, ["层级", "主要职责", "关键实现"], [
        ("界面层", "四行控制台、仪表、按钮、状态栏", "Qt Widgets / QSS / QPainter"),
        ("业务层", "状态机、控制映射、连接判定", "MainWindow + MavlinkManager"),
        ("协议层", "MAVLink打包、解析、ACK与参数", "ArduPilotMega MAVLink dialect"),
        ("传输层", "飞控数据双向传输", "QUdpSocket / UDP"),
        ("视频层", "RTSP解码、低延迟显示、录像", "FFmpeg"),
        ("输入层", "虚拟摇杆、键盘、游戏手柄", "Qt事件 + SDL2"),
        ("模型层", "STEP转换、网格加载、姿态绘制", "OpenCASCADE + QPainter"),
        ("伴随计算机", "串口桥接、UVC采集、RTSP发布", "Buildroot / socat / MediaMTX"),
    ], [1500, 3600, 4260], 8.8, True)

    add_heading(doc, "1.3 网络与端口基线", 2)
    add_table(doc, ["项目", "当前值", "说明"], [
        ("笔记本IP", "192.168.1.137/24", "有线网卡静态地址"),
        ("香橙派IP", "192.168.1.101/24", "网线或载波对端"),
        ("MAVLink目标", "192.168.1.101:15001/UDP", "Qt向香橙派发送"),
        ("Qt本地监听", "0.0.0.0:5555/UDP", "接收飞控返回数据"),
        ("RTSP", "192.168.1.101:554", "视频服务当前未自启动"),
        ("SSH", "192.168.1.101:22/TCP", "维护香橙派"),
    ], [2200, 2800, 4360], 9.0, True)

    page_break(doc)

    add_heading(doc, "2. 功能与当前状态", 1)
    add_table(doc, ["模块", "功能", "当前状态"], [
        ("MAVLink连接", "真实Pixhawk心跳判定、断线检测", "已实机验证"),
        ("遥测", "姿态、深度、速度、航向、电池、GPS、PWM", "已接入"),
        ("控制", "虚拟摇杆、实体手柄、键盘", "MANUAL_CONTROL 25Hz"),
        ("安全", "滑动解锁、滑动上锁、急停", "已接入MAVLink"),
        ("模式", "飞行模式、下潜/清洗姿态", "含固件定制接口"),
        ("云台", "前后云台方向与回中", "DO_SET_SERVO"),
        ("补光灯", "前后亮度滑动条", "PWM控制"),
        ("推进器", "6路功率条、PWM详情、反转/启用", "SERVO_OUTPUT_RAW + PARAM"),
        ("3D模型", "STEP导入、姿态跟随、360度观察", "已接入真实ATTITUDE"),
        ("视频", "双路RTSP、时间水印、重连", "Qt端完成；服务器自启动关闭"),
        ("录像", "AVI/MJPEG本地录像", "仅视频流就绪后可用"),
        ("日志", "姿态CSV、命令CSV", "程序运行时生成"),
    ], [1700, 4700, 2960], 8.8, True)

    add_heading(doc, "2.1 四行界面布局", 2)
    for item in (
        "第一行（20%）：连接状态、Pitch、Roll、Heading、推进器功率、帮助和设置。",
        "第二行（50%）：前摄像头、ROV 3D模型、深度仪表、后摄像头。",
        "第三行（30%）：补光灯、左右操作盘、前后云台、上锁/解锁和急停。",
        "第四行：Qt标准状态栏，显示解锁状态、系统状态、姿态和连接状态。",
    ):
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "2.2 已知边界", 2)
    for item in (
        "后摄像头需要香橙派实际发布 /live/main_stream_2，否则仅前路有画面。",
        "RTSP使用UDP低延迟模式，电力载波质量差时可能出现花屏或丢帧。",
        "MainWindow.cpp体量较大，新增复杂模块时应拆分为独立控制器。",
        "关闭流程保留线程terminate兜底；正常路径应先通过interrupt callback退出。",
        "深度公式以1013.25hPa为表面基准，现场高精度需求应增加零点标定。",
    ):
        add_bullet(doc, item, bullet_id)

    page_break(doc)

    add_heading(doc, "3. 通信协议与数据流", 1)
    add_callout(doc, "关键概念", "UDP是运输层，MAVLink是数据协议。程序不是在“UDP和MAVLink之间二选一”，而是在UDP数据报中传输MAVLink帧。")

    add_heading(doc, "3.1 接收消息", 2)
    add_table(doc, ["MAVLink消息", "用途", "界面/业务结果"], [
        ("HEARTBEAT", "在线、模式、解锁状态", "连接状态、模式、急停/安全判断"),
        ("SYS_STATUS", "电压、电流、电量、通信错误", "系统状态与电池信息"),
        ("ATTITUDE", "Roll/Pitch/Yaw及角速度", "三个仪表、3D模型、姿态日志"),
        ("VFR_HUD", "速度、航向、爬升率", "速度/深度速度信息"),
        ("SCALED_PRESSURE 1/2/3", "绝对压力和温度", "水深计算与深度仪表"),
        ("SERVO_OUTPUT_RAW", "舵机/推进器PWM", "6路功率条和PWM窗口"),
        ("RC_CHANNELS", "遥控通道", "输入监控"),
        ("GPS_RAW_INT", "定位与卫星", "GPS状态"),
        ("STATUSTEXT", "飞控文字提示", "状态栏/诊断信息"),
        ("COMMAND_ACK", "命令执行结果", "确认成功、拒绝或失败"),
        ("PARAM_VALUE", "参数读取结果", "参数/PID/PWM配置回显"),
    ], [2600, 2900, 3860], 8.6, True)

    add_heading(doc, "3.2 发送消息", 2)
    add_table(doc, ["命令/消息", "用途", "触发位置"], [
        ("GCS HEARTBEAT", "维持地面站心跳", "MavlinkManager定时器"),
        ("MANUAL_CONTROL", "ROV运动与Custom按键", "虚拟摇杆/手柄/键盘"),
        ("COMMAND_LONG ARM_DISARM", "解锁或上锁", "滑动开关/急停"),
        ("MAV_CMD_DO_SET_MODE", "切换飞行/清洗模式", "模式控件"),
        ("MAV_CMD_DO_SET_SERVO", "云台、灯光或舵机PWM", "方向按钮/滑动条"),
        ("PARAM_SET / REQUEST_READ", "修改并读取参数", "设置/PWM/PID窗口"),
        ("SET_MESSAGE_INTERVAL", "指定遥测消息频率", "首次真实心跳后"),
        ("SET_ATTITUDE_TARGET", "姿态目标控制", "设置窗口/清洗状态机"),
        ("COMMAND_LONG 31000", "固件定制一键清洗命令", "清洗模式逻辑"),
    ], [3000, 3300, 3060], 8.7, True)

    add_heading(doc, "3.3 Custom1/Custom2约定", 2)
    add_table(doc, ["功能", "MANUAL_CONTROL buttons", "固件配置", "当前行为"], [
        ("Custom1", "bit 1 / 数值2", "BTN1_FUNCTION=91", "角度 +5°"),
        ("Custom2", "bit 2 / 数值4", "BTN2_FUNCTION=92", "角度 -5°"),
    ], [1800, 2400, 2500, 2660], 9.0, True)
    add_body(doc, "若要更换手柄按键，修改上位机映射；若要把5°改为其他角度，原则上修改飞控固件功能实现。")

    page_break(doc)

    add_heading(doc, "4. 源码结构与修改入口", 1)
    add_table(doc, ["文件/目录", "职责", "常见修改"], [
        ("wz-724.pro", "qmake工程、依赖、编译后复制", "新增源文件、库、资源、DLL"),
        ("main.cpp", "程序入口与Qt事件循环", "启动参数、调试自动退出"),
        ("MainWindow.ui", "Designer基础控件", "新增静态控件和objectName"),
        ("MainWindow.hpp/.cpp", "主窗口和业务总调度", "按钮、布局、状态、控制映射"),
        ("mavlink/mavlink_manager.*", "MAVLink状态、解析和发送", "新消息、新命令、新参数"),
        ("udphandler.*", "QUdpSocket封装", "本地绑定、目标地址、错误处理"),
        ("ffmpegrtspplayer.*", "RTSP解码和录像", "延迟、编码、线程退出"),
        ("attitude3dwidget.*", "网格加载和3D绘制", "模型方向、视角、交互"),
        ("rovdashboardwidgets.*", "自绘仪表", "Pitch/Roll/Heading/深度/PWM外观"),
        ("virtualjoystick.*", "虚拟双摇杆", "量程、回中和发送节奏"),
        ("sdl2gamepad.*", "实体手柄采集", "手柄轴和按钮"),
        ("pwmdisplaywidget.*", "PWM、反转、启用、PID", "参数映射与显示"),
        ("styles/professional.qss", "专业深色主题", "颜色、字体、边框、悬停"),
        ("models/ + models.qrc", "默认3D网格资源", "替换内置模型"),
        ("tools/step_converter.py", "STEP转wzmesh", "三角化精度和颜色"),
        ("orangepi-deployment/", "伴随计算机部署模板", "桥接、视频服务配置"),
        ("tests/", "协议回归与实机探针", "新增协议测试用例"),
    ], [2650, 3300, 3410], 8.4, True)

    add_heading(doc, "4.1 不要直接修改的文件", 2)
    for item in (
        "ui_MainWindow.h、ui_dialogsetup.h：由uic根据.ui生成，重新构建会覆盖。",
        "thirdparty/mavlink生成头文件：除非重新生成dialect，否则不要手工改。",
        "便携目录中的DLL：应从对应Qt/FFmpeg/SDL2版本重新部署，不要随意替换单个文件。",
        ".pro.user：仅保存本机Qt Creator Kit，不应作为源码配置交接。",
    ):
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "4.2 功能修改定位表", 2)
    add_table(doc, ["需求", "优先修改位置"], [
        ("改变四行比例/控件位置", "MainWindow::setupFourRowDashboard()"),
        ("改变颜色/圆角/字体", "styles/professional.qss"),
        ("改变仪表绘制", "rovdashboardwidgets.cpp"),
        ("增加MAVLink接收消息", "parseMavlinkData() + 新handle函数 + VehicleState"),
        ("增加飞控控制命令", "MavlinkManager新增send函数，MainWindow只负责触发"),
        ("调整摇杆方向/死区", "onJoystickValueChanged()/onManualControlTimer()"),
        ("调整Custom1/2按钮", "MainWindow顶部掩码 + connectGamepadSignals()"),
        ("调整3D Roll/Pitch/Yaw方向", "Attitude3DWidget::paintEvent()中的rotate"),
        ("调整3D默认视角", "Attitude3DWidget::resetView()"),
        ("调整RTSP地址", "frontCameraUrl()/rearCameraUrl()"),
        ("调整录像格式/码率", "FFmpegRtspPlayer::initializeRecorder()"),
        ("处理关闭卡死", "stopBackgroundServices()/FFmpegRtspPlayer::shutdown()"),
        ("改变USB摄像头采集参数", "香橙派MediaMTX/FFmpeg配置，不在Qt代码"),
        ("改变固件+5°动作本身", "飞控固件BTN1/BTN2功能实现"),
    ], [3900, 5460], 8.7, True)

    page_break(doc)

    add_heading(doc, "5. 开发环境与构建发布", 1)
    add_heading(doc, "5.1 推荐开发环境", 2)
    add_table(doc, ["项目", "建议版本/要求"], [
        ("操作系统", "Windows 10/11 64位"),
        ("Qt", "Qt 5.14.2"),
        ("编译器", "MinGW 64-bit，必须与Qt Kit一致"),
        ("构建系统", "qmake / .pro"),
        ("语言标准", "C++11"),
        ("外部库", "FFmpeg 62系列、SDL2 2.30.10、MAVLink头文件"),
    ], [2800, 6560], 9.2, True)

    add_heading(doc, "5.2 Qt Creator首次导入", 2)
    for item in (
        "备份整个工程目录。",
        "若出现No valid settings file，关闭Qt Creator并移走wz-724.pro.user。",
        "用Qt Creator打开wz-724.pro，而不是打开旧的.user文件。",
        "选择Desktop Qt 5.14.2 MinGW 64-bit Kit。",
        "执行“构建 > 运行qmake”。",
        "执行“清理项目”，然后“重新构建项目”。",
        "先运行Debug版本确认功能，再生成Release版本。",
    ):
        add_number(doc, item, decimal_id)

    add_heading(doc, "5.3 发布注意事项", 2)
    for item in (
        "不能只复制wz-724.exe；目标电脑还需要Qt、平台插件、FFmpeg、SDL2及MinGW运行库。",
        "发布时以已验证的wz-724-portable-v17-mavlink目录为基线。",
        "wz-step-converter.exe必须与主程序同目录，否则STEP导入不可用。",
        "修改源码后必须重新构建，再替换便携目录中的exe；新增依赖时重新部署DLL。",
        "压缩包交付后，应在未安装Qt的电脑上做一次启动和关闭测试。",
    ):
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "5.4 回归测试", 2)
    add_code_block(doc, [
        "# MAVLink协议单元测试（Qt Creator打开测试pro或运行构建产物）",
        "tests/mavlink_manager_protocol_test.pro",
        "",
        "# 实机只发送GCS心跳并接收遥测，不发送解锁/运动命令",
        "python tests/mavlink_udp_probe.py --host 192.168.1.101 --remote-port 15001 --local-port 5555",
    ])

    page_break(doc)

    add_heading(doc, "6. 香橙派部署与维护", 1)
    add_heading(doc, "6.1 MAVLink桥接", 2)
    add_body(doc, "现场实测桥接服务为 /etc/init.d/S100_mavlink-bridge，使用 /dev/ttyS0 @ 57600，并在UDP 15001监听。Qt从本地5555发出GCS心跳后，socat建立双向会话。")
    add_code_block(doc, [
        "# 查看服务与端口",
        "/etc/init.d/S100_mavlink-bridge status",
        "ps -ef | grep socat",
        "stty -F /dev/ttyS0 -a",
        "netstat -an | grep 15001",
        "tail -n 100 /var/log/socat_mavlink.log",
    ])
    add_callout(doc, "禁止并行桥接", "不要同时启动S100_mavlink-bridge和项目模板中的S94wz-mavlink，否则会争用/dev/ttyS0和UDP 15001。", PALE_RED, "9B1C1C")

    add_heading(doc, "6.2 摄像头与RTSP", 2)
    add_body(doc, "USB摄像头正常路径为UVC/V4L2设备 /dev/video0，MediaMTX按配置启动FFmpeg发布RTSP。Qt上位机只接收RTSP，不直接打开UVC设备。")
    add_table(doc, ["项目", "当前状态"], [
        ("原启动脚本", "/etc/init.d/S95wz-camera"),
        ("当前禁用文件", "/etc/init.d/wz-camera.disabled"),
        ("当前自启动", "已关闭"),
        ("前路地址", "rtsp://192.168.1.101/live/main_stream_1"),
        ("后路地址", "rtsp://192.168.1.101/live/main_stream_2"),
    ], [3000, 6360], 9.2, True)
    add_code_block(doc, [
        "# 恢复摄像头自启动并立即启动",
        "mv /etc/init.d/wz-camera.disabled /etc/init.d/S95wz-camera",
        "chmod 755 /etc/init.d/S95wz-camera",
        "/etc/init.d/S95wz-camera start",
        "",
        "# 再次关闭",
        "/etc/init.d/S95wz-camera stop",
        "mv /etc/init.d/S95wz-camera /etc/init.d/wz-camera.disabled",
        "chmod 644 /etc/init.d/wz-camera.disabled",
    ])

    add_heading(doc, "6.3 常用只读诊断", 2)
    add_code_block(doc, [
        "ip addr",
        "lsusb",
        "ls -l /dev/video* /dev/ttyS* 2>/dev/null",
        "ffmpeg -hide_banner -f v4l2 -list_formats all -i /dev/video0",
        "ps -ef | grep -E 'socat|ffmpeg|mediamtx'",
        "netstat -an",
        "dmesg | tail -n 100",
    ])

    page_break(doc)

    add_heading(doc, "7. 运行与现场测试SOP", 1)
    add_heading(doc, "7.1 上电前", 2)
    for item in (
        "推进器处于安全状态，首次测试断动力或拆桨。",
        "确认急停按钮可操作，滑动解锁处于未触发状态。",
        "检查Pixhawk、香橙派、摄像头和网线/载波连接。",
        "笔记本有线网卡设置为192.168.1.137/24，不配置错误网关。",
        "关闭可能占用UDP 5555的QGC或旧版wz-724。",
    ):
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "7.2 联机步骤", 2)
    for item in (
        "ping 192.168.1.101，确认链路无丢包。",
        "启动wz-724便携版，等待真实Pixhawk心跳。",
        "确认连接状态显示在线，而不是只显示UDP已绑定。",
        "缓慢转动飞控，核对Pitch、Roll、Heading和3D模型方向。",
        "核对深度、速度、电池和6路PWM显示是否合理。",
        "视频服务启用时，核对前后画面、时间水印和录像按钮。",
        "保持未解锁状态测试云台、Custom1/2和补光灯。",
        "动力测试必须由现场安全负责人确认后单独进行。",
    ):
        add_number(doc, item, decimal_id)

    add_heading(doc, "7.3 验收记录", 2)
    add_table(doc, ["检查项", "预期结果", "结果/签字"], [
        ("网络", "香橙派可ping通，SSH可连接", ""),
        ("MAVLink", "真实sysid=1心跳，状态在线", ""),
        ("姿态", "仪表与3D方向正确、无跳变", ""),
        ("深度", "表面附近合理，可随压力变化", ""),
        ("手柄", "轴、死区、Custom1/2正确", ""),
        ("安全", "解锁需确认，急停立即上锁", ""),
        ("视频", "启用服务后画面和水印正常", ""),
        ("录像", "AVI可播放，文件正常结束", ""),
        ("退出", "窗口关闭后进程和线程退出", ""),
    ], [2500, 4700, 2160], 8.8, True)

    page_break(doc)

    add_heading(doc, "8. 常见故障与排查", 1)
    add_table(doc, ["现象", "优先检查", "处理建议"], [
        ("显示未连接", "ping、5555占用、15001、socat、串口", "先用UDP探针确认真实HEARTBEAT"),
        ("只绑定UDP但不上线", "是否收到自动驾驶仪心跳", "不要把GCS心跳当Pixhawk心跳"),
        ("姿态方向反了", "3D坐标系与Pixhawk NED符号", "调整paintEvent中的rotate正负号"),
        ("深度一直为0", "SCALED_PRESSURE1/2/3是否到达", "核对传感器、消息频率和表面基准"),
        ("无视频", "摄像头服务当前是否禁用", "恢复wz-camera服务并检查/dev/video0"),
        ("视频卡顿", "载波丢包、码率、RTSP UDP", "降到720p及2-4Mbps，检查链路"),
        ("录像按钮不可用", "视频流和videoSize是否就绪", "先等待首帧，再检查写入目录"),
        ("关闭软件卡死", "RTSP线程/异步连接/重连定时器", "检查shutdown和interrupt callback"),
        ("STEP导入失败", "转换器是否同目录、文件是否有效", "直接运行转换器查看错误输出"),
        ("手柄按钮错位", "SDL映射和gamecontrollerdb", "记录GUID并更新映射"),
        ("Custom1/2无动作", "buttons值、固件BTN功能、当前模式", "确认值2/4且固件功能91/92"),
        ("其他电脑打不开exe", "Qt/FFmpeg/SDL2 DLL与平台插件", "发送完整便携目录而非单exe"),
    ], [2200, 3500, 3660], 8.5, True)

    add_heading(doc, "8.1 Windows端常用命令", 2)
    add_code_block(doc, [
        "ping 192.168.1.101",
        "ipconfig /all",
        "arp -a",
        "netstat -ano -p udp | findstr :5555",
        "tasklist | findstr wz-724",
    ])

    add_heading(doc, "8.2 日志位置", 2)
    add_table(doc, ["日志", "用途"], [
        ("log/*_attitude.csv", "姿态、角速度和累计Yaw分析"),
        ("log/*_command.csv", "MAVLink命令、参数和buttons追踪"),
        ("/var/log/socat_mavlink.log", "香橙派串口桥接连接记录"),
        ("/var/log/wz-mediamtx.log", "RTSP服务和FFmpeg发布错误"),
    ], [3300, 6060], 9.0, True)

    page_break(doc)

    add_heading(doc, "9. 变更开发规范", 1)
    add_heading(doc, "9.1 每次改功能的推荐流程", 2)
    for item in (
        "明确边界：改动属于Qt界面、MAVLink、香橙派视频，还是飞控固件。",
        "先建立可回退基线：备份或Git提交。",
        "用Qt Creator全局搜索objectName、槽函数或MAVLink消息名。",
        "界面只负责触发和显示，协议打包集中在MavlinkManager。",
        "新增接收消息时，同时更新VehicleState、信号和界面刷新。",
        "新增控制命令时，记录目标sysid/compid、参数含义和ACK结果。",
        "先运行协议测试，再做未解锁实机测试。",
        "检查启动、断线重连、录像、退出和无设备场景。",
        "重新生成Release便携版，并在无Qt电脑上冒烟测试。",
        "更新本交接文档的版本、变更记录和已知问题。",
    ):
        add_number(doc, item, decimal_id)

    add_heading(doc, "9.2 编码原则", 2)
    for item in (
        "不要在多个按钮槽中重复拼MAVLink消息，统一封装send函数。",
        "不要把UDP绑定成功当作飞控在线，必须基于真实自动驾驶仪HEARTBEAT。",
        "不要在UI线程中执行阻塞式RTSP读取、长时间sleep或STEP转换。",
        "新增后台线程必须设计停止标志、超时和析构顺序。",
        "修改PWM、解锁、模式和参数写入必须保留安全条件和ACK检查。",
        "避免硬编码新增IP/端口，优先加入设置并保存到QSettings。",
    ):
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "9.3 建议的后续重构", 2)
    add_table(doc, ["优先级", "建议", "收益"], [
        ("高", "初始化Git并建立发布标签", "可追溯、可回滚"),
        ("高", "拆分MainWindow为Telemetry/Control/Video控制器", "降低耦合和回归风险"),
        ("高", "增加硬件在环自动验收脚本", "减少现场人工判断"),
        ("中", "把网络和RTSP配置集中成一份结构化配置", "避免端口散落"),
        ("中", "增加深度零点校准", "提高现场深度准确性"),
        ("中", "为视频增加TCP/UDP可切换和码率档位", "适应载波链路"),
        ("低", "将软件3D渲染替换为Qt3D/OpenGL", "大型模型性能更好"),
    ], [1300, 5000, 3060], 8.8, True)

    page_break(doc)

    add_heading(doc, "10. 安全要求与交接签署", 1)
    add_callout(doc, "安全红线", "任何涉及解锁、推进器、舵机极限、模式切换或参数写入的测试，都必须先在无动力或物理隔离状态完成。严禁仅依据界面显示判断设备安全。", PALE_RED, "9B1C1C")

    add_heading(doc, "10.1 接手确认清单", 2)
    add_table(doc, ["确认内容", "完成"], [
        ("已收到源码、依赖、便携版、测试和部署模板", "□"),
        ("已能用Qt Creator重新qmake并完成Release构建", "□"),
        ("已理解UDP与MAVLink的分层关系", "□"),
        ("已能确认真实Pixhawk心跳而非GCS心跳", "□"),
        ("已完成未解锁姿态、深度、手柄和急停测试", "□"),
        ("已知摄像头自启动当前关闭及恢复方法", "□"),
        ("已了解Custom1/2和固件定制命令", "□"),
        ("已建立备份或版本控制基线", "□"),
    ], [8200, 1160], 9.0, True)

    add_heading(doc, "10.2 交接记录", 2)
    add_table(doc, ["角色", "姓名/签字", "日期", "备注"], [
        ("移交人", "", "", ""),
        ("接收人", "", "", ""),
        ("现场安全确认", "", "", ""),
    ], [1700, 2700, 1900, 3060], 9.2, True)

    add_heading(doc, "附录A：关键参数速查", 1)
    add_table(doc, ["参数", "值"], [
        ("Qt目标名称", "wz-724"),
        ("Qt版本", "5.14.2"),
        ("编译器", "MinGW 64-bit"),
        ("香橙派", "192.168.1.101"),
        ("笔记本", "192.168.1.137"),
        ("MAVLink远端", "UDP 15001"),
        ("Qt本地", "UDP 5555"),
        ("飞控串口", "/dev/ttyS0 @ 57600"),
        ("前视频", "/live/main_stream_1"),
        ("后视频", "/live/main_stream_2"),
        ("Custom1", "buttons=2 / BTN1_FUNCTION=91"),
        ("Custom2", "buttons=4 / BTN2_FUNCTION=92"),
        ("录像", "AVI + MJPEG"),
    ], [3600, 5760], 9.2, True)

    # Apply keep-together rules and explicit font to all body runs.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            if run.font.size is None:
                set_run_font(run, size=10.5, color=INK)

    doc.core_properties.title = "WZ-724上位机工作交接文档"
    doc.core_properties.subject = "ROV上位机开发、部署、运行与维护交接"
    doc.core_properties.author = "WZ-724 ROV项目组"
    doc.core_properties.keywords = "WZ-724, ROV, Qt, MAVLink, Pixhawk, Orange Pi, RTSP"
    doc.core_properties.comments = "技术交接基线 V1.0"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
